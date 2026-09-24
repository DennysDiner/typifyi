"""Text extraction: PDF (pypdf), DOCX (python-docx), HTML (bs4), plain text; OCR fallback via ocrmypdf
when a PDF yields little text and ocrmypdf/tesseract are installed. Text is stored in document_text
(FTS5-indexed) next to the immutable original."""
from __future__ import annotations

import io
import shutil
import sqlite3
import subprocess
import tempfile
from pathlib import Path

from .db import tx, utcnow

MIN_CHARS_PER_PAGE = 40


def pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for p in reader.pages:
        try:
            parts.append(p.extract_text() or "")
        except Exception:  # noqa: BLE001
            parts.append("")
    return "\n\f".join(parts), len(reader.pages)


def docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def html_text(data: bytes) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "lxml")
    for t in soup(["script", "style", "nav", "header", "footer"]):
        t.decompose()
    main = soup.select_one("main") or soup.body or soup
    return main.get_text("\n", strip=True)


def ocr_available() -> bool:
    return shutil.which("ocrmypdf") is not None and shutil.which("tesseract") is not None


def ocr_pdf(data: bytes, timeout: int = 900) -> bytes | None:
    if not ocr_available():
        return None
    with tempfile.TemporaryDirectory() as td:
        src, dst = Path(td) / "in.pdf", Path(td) / "out.pdf"
        src.write_bytes(data)
        try:
            subprocess.run(["ocrmypdf", "--skip-text", "--quiet", "-l", "eng", str(src), str(dst)], check=True, timeout=timeout,
                           capture_output=True)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            return None
        return dst.read_bytes() if dst.exists() else None


def extract(data: bytes, content_type: str | None, filename: str | None = None, allow_ocr: bool = True) -> dict:
    ct = (content_type or "").split(";")[0].strip().lower()
    name = (filename or "").lower()
    if data[:5] == b"%PDF-" or ct == "application/pdf" or name.endswith(".pdf"):
        text, pages = pdf_text(data)
        method, ocr = "pypdf", 0
        if allow_ocr and pages and len(text.strip()) < MIN_CHARS_PER_PAGE * pages:
            ocred = ocr_pdf(data)
            if ocred:
                text2, _ = pdf_text(ocred)
                if len(text2.strip()) > len(text.strip()):
                    text, method, ocr = text2, "ocrmypdf+pypdf", 1
            elif not ocr_available():
                method = "pypdf (OCR needed but ocrmypdf/tesseract not installed)"
        return {"text": text, "pages": pages, "method": method, "ocr": ocr}
    if ct.endswith("wordprocessingml.document") or name.endswith(".docx"):
        return {"text": docx_text(data), "pages": None, "method": "python-docx", "ocr": 0}
    if ct in ("text/html", "application/xhtml+xml") or name.endswith((".html", ".htm")):
        return {"text": html_text(data), "pages": None, "method": "bs4", "ocr": 0}
    if ct.startswith("text/") or name.endswith((".txt", ".csv", ".md")):
        return {"text": data.decode("utf-8", errors="replace"), "pages": None, "method": "text", "ocr": 0}
    return {"text": "", "pages": None, "method": f"unsupported:{ct or name}", "ocr": 0}


def extract_pending(conn: sqlite3.Connection, archive, limit: int = 50, allow_ocr: bool = True) -> int:
    rows = conn.execute(
        "SELECT d.id, d.sha256, d.content_type, d.filename FROM documents d WHERE d.text_extracted_at IS NULL ORDER BY d.id LIMIT ?", (limit,)
    ).fetchall()
    n = 0
    for r in rows:
        data = archive.read(r["sha256"])
        res = extract(data, r["content_type"], r["filename"], allow_ocr=allow_ocr)
        with tx(conn):
            conn.execute("INSERT OR REPLACE INTO document_text(document_id,text) VALUES (?,?)", (r["id"], res["text"]))
            conn.execute(
                "UPDATE documents SET text_extracted_at=?, extraction_method=?, ocr_applied=?, page_count=?, text_chars=? WHERE id=?",
                (utcnow(), res["method"], res["ocr"], res["pages"], len(res["text"]), r["id"]),
            )
        n += 1
    return n
